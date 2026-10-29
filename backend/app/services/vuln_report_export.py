"""Vulnerability report export — solution-set CSV format (Risk, Host, Protocol, Port, Name, Synopsis, Description, Solution)."""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Any

from app.db.sql_helpers import execute, fetchall, fetchone
from app.services.aetheris_severity import aetheris_severity_sql, classify_severity
from app.services.vuln_correlation import correlation_key_for
from app.services.vuln_finding_ingest import refresh_asset_risk, upsert_finding

CSV_COLUMNS = ("Risk", "Host", "Protocol", "Port", "Name", "Synopsis", "Description", "Solution")
# Matches the user-approved "solution set - new.xlsx" export: eight columns.
XLSX_COLUMNS = ("Risk", "Host", "Protocol", "Port", "Name", "Synopsis", "Description", "Solution")
XLSX_COL_WIDTHS = {
    "A": 8.88671875,
    "B": 15.88671875,
    "C": 8.88671875,
    "D": 13.0,
    "E": 60.6640625,
    "F": 62.88671875,
    "G": 8.88671875,
    "H": 13.0,
}
# Same fills as Solution Set (2).xlsx Risk column.
XLSX_RISK_FILLS = {
    "Critical": "FFC00000",
    "High": "FFFF0000",
    "Medium": "FFFFC000",
    "Low": "FF92D050",
    "Info": "FF5B9BD5",
}

_SEV_TO_RISK = {
    "critical": "Critical",
    "high": "High",
    "medium": "Medium",
    "low": "Low",
    "info": "Info",
}


def job_kw(scan_job_id: str | None) -> dict[str, str]:
    """``{"scan_job_id": id}`` when a job is selected, else ``{}`` (V1.10.3).

    Passing the filter only when it is set keeps every existing case-level
    caller (and its call signature) unchanged.
    """
    jid = str(scan_job_id or "").strip()
    return {"scan_job_id": jid} if jid else {}


def _job_clause(scan_job_id: str | None, params: dict[str, Any], *, column: str = "f.scan_job_id") -> str:
    jid = str(scan_job_id or "").strip()
    if not jid:
        return ""
    params["jid"] = jid
    return f" AND {column} = CAST(:jid AS uuid)"


def _risk_label(severity: str | None, cvss: float | None, synopsis: str = "") -> str:
    sev, _ = classify_severity(cvss=float(cvss or 0), synopsis=synopsis, raw_severity=severity)
    return _SEV_TO_RISK.get(sev, "Info")


def _open_finding_rows(db, *, case_id: str, scan_job_id: str | None = None) -> list[dict[str, Any]]:
    params: dict[str, Any] = {"cid": case_id}
    job_sql = _job_clause(scan_job_id, params)
    return fetchall(
        db,
        f"""SELECT f.severity, f.cvss, f.synopsis, f.description, f.remediation, f.port, f.protocol,
                  f.scan_engine, f.cve, f.correlation_key,
                  COALESCE(a.primary_ip, a.hostname, '') AS host
           FROM vuln_findings f
           LEFT JOIN vuln_assets a ON a.id = f.asset_id
           WHERE f.case_id = CAST(:cid AS uuid) AND f.status = 'open'{job_sql}
           ORDER BY
             CASE LOWER(COALESCE(f.severity, ''))
               WHEN 'critical' THEN 1
               WHEN 'high' THEN 2
               WHEN 'medium' THEN 3
               WHEN 'low' THEN 4
               WHEN 'info' THEN 5
               ELSE 6
             END,
             COALESCE(a.primary_ip, a.hostname, ''),
             COALESCE(f.port, 0),
             COALESCE(f.synopsis, '')""",
        params,
    )


def _solution_row(r: dict[str, Any]) -> dict[str, Any]:
    port = r.get("port")
    try:
        port_val = int(port) if port is not None and str(port).strip() != "" else 0
    except (TypeError, ValueError):
        port_val = 0
    return {
        "Risk": _risk_label(r.get("severity"), r.get("cvss"), str(r.get("synopsis") or "")),
        "Host": r.get("host") or "",
        "Protocol": (r.get("protocol") or "tcp") or "tcp",
        "Port": port_val,
        "Name": r.get("synopsis") or r.get("cve") or "Finding",
        "Synopsis": (r.get("synopsis") or "")[:500],
        "Description": r.get("description") or "",
        "Solution": r.get("remediation") or "",
    }


def export_case_csv(db, *, case_id: str, scan_job_id: str | None = None) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=CSV_COLUMNS, lineterminator="\n")
    writer.writeheader()
    for r in _open_finding_rows(db, case_id=case_id, **job_kw(scan_job_id)):
        writer.writerow(_solution_row(r))
    return buf.getvalue()


def write_solution_xlsx(rows: list[dict[str, Any]], *, widen_host: bool = False) -> bytes:
    """Write the approved Solution Set workbook with 8 columns and matching layout."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    header_font = Font(name="Aptos Narrow", size=11, bold=False)
    body_font = Font(name="Aptos Narrow", size=11)
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    risk_fills = {
        risk: PatternFill(start_color=color, end_color=color, fill_type="solid")
        for risk, color in XLSX_RISK_FILLS.items()
    }
    for col_idx, title in enumerate(XLSX_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.border = thin
        cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    for r_idx, row in enumerate(rows, start=2):
        for col_idx, key in enumerate(XLSX_COLUMNS, start=1):
            cell = ws.cell(row=r_idx, column=col_idx, value=row[key])
            cell.font = body_font
            cell.border = thin
            if key == "Risk":
                fill = risk_fills.get(str(row[key]))
                if fill is not None:
                    cell.fill = fill
            if key in {"Synopsis", "Description", "Solution"}:
                cell.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[r_idx].height = 13.95
    for col, width in XLSX_COL_WIDTHS.items():
        ws.column_dimensions[col].width = width
    if widen_host and rows:
        longest = max(len(str(row.get("Host") or "")) for row in rows)
        ws.column_dimensions["B"].width = min(80.0, max(XLSX_COL_WIDTHS["B"], float(longest + 2)))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_case_xlsx(db, *, case_id: str, scan_job_id: str | None = None) -> bytes:
    """Write the approved Solution Set workbook with 8 columns and matching layout."""
    rows = [_solution_row(raw) for raw in _open_finding_rows(db, case_id=case_id, **job_kw(scan_job_id))]
    return write_solution_xlsx(rows)


_SOLUTION_RISK_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "Info": 4}


def _norm_solution_text(val: Any) -> str:
    return re.sub(r"\s+", " ", str(val if val is not None else "").strip().lower())


def _solution_merge_key(row: dict[str, Any]) -> tuple[str, str, str, str, str, str]:
    """Same identity as a merged gap finding: everything except the host."""
    return (
        _norm_solution_text(row.get("Name")),
        _norm_solution_text(row.get("Risk")),
        _norm_solution_text(row.get("Description")),
        _norm_solution_text(row.get("Solution")),
        _norm_solution_text(row.get("Port")),
        _norm_solution_text(row.get("Protocol")),
    )


def _host_parts(host: Any) -> list[str]:
    return [part.strip() for part in str(host or "").split(",") if part.strip()]


def merge_solution_rows(row_lists: list[list[dict[str, Any]]]) -> list[dict[str, Any]]:
    """Combine Solution Set rows the way gap reports combine findings.

    Rows that share name, risk, description, solution, port and protocol become
    one row. Host addresses are appended, comma-separated, each address once.
    A finding that differs on any of those fields stays on its own row.
    """
    merged: dict[tuple[str, str, str, str, str, str], dict[str, Any]] = {}
    order: list[tuple[str, str, str, str, str, str]] = []
    for rows in row_lists:
        for row in rows:
            if not isinstance(row, dict):
                continue
            key = _solution_merge_key(row)
            hosts = _host_parts(row.get("Host"))
            if key not in merged:
                item = dict(row)
                item["_hosts"] = hosts
                merged[key] = item
                order.append(key)
                continue
            existing = merged[key]
            seen = {host.lower() for host in existing["_hosts"]}
            for host in hosts:
                if host.lower() not in seen:
                    seen.add(host.lower())
                    existing["_hosts"].append(host)
    out: list[dict[str, Any]] = []
    for key in order:
        item = merged[key]
        hosts = sorted(item.pop("_hosts"), key=str.lower)
        item["Host"] = ", ".join(hosts)
        out.append(item)
    out.sort(
        key=lambda row: (
            _SOLUTION_RISK_ORDER.get(str(row.get("Risk") or ""), 5),
            str(row.get("Name") or "").lower(),
            str(row.get("Host") or "").lower(),
        )
    )
    return out


def _ordered_case_ids(case_ids: list[str] | None) -> list[str]:
    ordered: list[str] = []
    for raw in case_ids or []:
        cid = str(raw or "").strip()
        if cid and cid not in ordered:
            ordered.append(cid)
    if len(ordered) < 2:
        raise ValueError("Select at least two cases to merge solution set reports")
    return ordered


def export_merged_solution_xlsx(
    db,
    case_ids: list[str] | None = None,
    *,
    job_ids: list[str] | None = None,
) -> tuple[bytes, str, str]:
    """One Solution Set workbook for the same jobs (or cases) a merged gap PDF uses.

    Each part is the workbook that job or case would download on its own. Identical
    findings are combined; their hosts are listed together in the Host column.
    """
    lists: list[list[dict[str, Any]]] = []
    if job_ids:
        from app.services.gap_assessment_report import load_merge_job_refs

        refs = load_merge_job_refs(db, list(job_ids))
        for ref in refs:
            raw_rows = _open_finding_rows(db, case_id=str(ref["case_id"]), scan_job_id=str(ref["id"]))
            lists.append([_solution_row(raw) for raw in raw_rows])
        filename = f"Solution Set-merged-{len(refs)}-jobs.xlsx"
    else:
        ids = _ordered_case_ids(case_ids)
        for cid in ids:
            raw_rows = _open_finding_rows(db, case_id=cid)
            lists.append([_solution_row(raw) for raw in raw_rows])
        filename = f"Solution Set-merged-{len(ids)}-cases.xlsx"
    data = write_solution_xlsx(merge_solution_rows(lists), widen_host=True)
    return (
        data,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename,
    )


def export_case_summary(db, *, case_id: str, scan_job_id: str | None = None) -> dict[str, Any]:
    params: dict[str, Any] = {"cid": case_id}
    job_sql = _job_clause(scan_job_id, params, column="scan_job_id")
    rows = fetchall(
        db,
        f"""SELECT {aetheris_severity_sql('')} AS severity, COUNT(*) AS c FROM vuln_findings
           WHERE case_id = CAST(:cid AS uuid) AND status = 'open'{job_sql}
           GROUP BY 1""",
        params,
    )
    by_sev = {str(r["severity"]): int(r["c"]) for r in rows}
    engines = fetchall(
        db,
        f"""SELECT scan_engine, COUNT(*) AS c FROM vuln_findings
           WHERE case_id = CAST(:cid AS uuid) AND status = 'open' AND scan_engine IS NOT NULL{job_sql}
           GROUP BY scan_engine ORDER BY c DESC""",
        params,
    )
    if scan_job_id:
        # Correlation groups are built across the whole case; count only groups
        # that contain at least one open finding of this job.
        corr = fetchone(
            db,
            """SELECT COUNT(*) AS total, SUM(CASE WHEN c.engine_count > 1 THEN 1 ELSE 0 END) AS multi
               FROM vuln_finding_correlations c
               WHERE c.case_id = CAST(:cid AS uuid)
                 AND EXISTS (SELECT 1 FROM vuln_findings f
                             WHERE f.case_id = c.case_id AND f.status = 'open'
                               AND f.correlation_key = c.correlation_key
                               AND f.scan_job_id = CAST(:jid AS uuid))""",
            params,
        )
    else:
        corr = fetchone(
            db,
            """SELECT COUNT(*) AS total, SUM(CASE WHEN engine_count > 1 THEN 1 ELSE 0 END) AS multi
               FROM vuln_finding_correlations WHERE case_id = CAST(:cid AS uuid)""",
            {"cid": case_id},
        )
    return {
        "case_id": case_id,
        "scan_job_id": str(scan_job_id) if scan_job_id else None,
        "findings_by_severity": by_sev,
        "findings_by_engine": {str(r["scan_engine"]): int(r["c"]) for r in engines},
        "correlation_groups": int((corr or {}).get("total") or 0),
        "multi_engine_groups": int((corr or {}).get("multi") or 0),
        "service_coverage": load_service_coverage(db, case_id=case_id, **job_kw(scan_job_id)),
    }


def load_service_coverage(db, *, case_id: str, scan_job_id: str | None = None) -> list[dict[str, Any]]:
    """Export durable engine/family evidence separately from vulnerability rows."""
    params: dict[str, Any] = {"cid": case_id}
    job_sql = _job_clause(scan_job_id, params, column="id")
    rows = fetchall(db, f"""SELECT id, status, orchestration_json FROM vuln_scan_jobs
        WHERE case_id = CAST(:cid AS uuid){job_sql} ORDER BY created_at""", params)
    out = []
    for row in rows:
        orch = row.get("orchestration_json") or {}
        if isinstance(orch, str):
            try:
                orch = json.loads(orch)
            except ValueError:
                continue
        if not isinstance(orch, dict) or not orch.get("edge_agent"):
            continue
        out.append({"job_id": str(row.get("id") or ""), "status": row.get("status"),
                    "service_coverage": orch.get("service_coverage") or {},
                    "policy_snapshots": orch.get("policy_snapshots") or {},
                    "limitation": "Missing native, authenticated or network-scoped evidence is not a per-rule pass"})
    return out


def export_service_coverage_csv(db, *, case_id: str, scan_job_id: str | None = None) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(["Job", "Host", "Service ID", "Service", "Evidence state", "Reason", "Source result IDs", "Credentialed coverage"])
    for job in load_service_coverage(db, case_id=case_id, **job_kw(scan_job_id)):
        for host, evidence in job["service_coverage"].items():
            for service in evidence.get("services") or []:
                writer.writerow([job["job_id"], host, service.get("service_id"), service.get("name"),
                                 service.get("status"), service.get("reason"),
                                 ",".join(service.get("source_result_ids") or []), service.get("credentialed_checks")])
            forwarding = evidence.get("ip_forwarding") or {}
            writer.writerow([job["job_id"], host, "CFG-IP-FORWARDING", "IP forwarding", forwarding.get("status", "not-tested"),
                             forwarding.get("reason", "No authenticated configuration evidence"), forwarding.get("source_ref", ""), "Requires device-role policy"])
    return buf.getvalue()


def import_solution_csv(db, *, case_id: str, csv_text: str, scan_job_id: str | None = None) -> dict[str, Any]:
    """Import Nessus/OpenVAS-style CSV (solution set format) into vuln_findings."""
    reader = csv.DictReader(io.StringIO(csv_text))
    if not reader.fieldnames:
        return {"imported": 0, "error": "empty csv"}
    norm_fields = {f.strip().lower(): f for f in reader.fieldnames if f}
    required = {"risk", "host", "name"}
    if not required.issubset(norm_fields.keys()):
        return {"imported": 0, "error": f"missing columns; need {required}"}

    imported = 0
    for row in reader:
        host = (row.get(norm_fields["host"]) or "").strip()
        if not host:
            continue
        name = row.get(norm_fields.get("name", "Name")) or ""
        risk = (row.get(norm_fields.get("risk", "Risk")) or "Medium").strip()
        sev_map = {"critical": "critical", "high": "high", "medium": "medium", "low": "low", "info": "info"}
        severity = sev_map.get(risk.lower(), "medium")
        port_raw = row.get(norm_fields.get("port", "Port")) if "port" in norm_fields else None
        port = int(port_raw) if port_raw and str(port_raw).strip().isdigit() else None
        protocol = (row.get(norm_fields.get("protocol", "Protocol")) or "tcp").strip() if "protocol" in norm_fields else "tcp"
        synopsis = row.get(norm_fields.get("synopsis", "Synopsis")) or name
        description = row.get(norm_fields.get("description", "Description")) or ""
        solution = row.get(norm_fields.get("solution", "Solution")) or ""

        asset = fetchone(
            db,
            """SELECT id FROM vuln_assets
               WHERE case_id = CAST(:cid AS uuid) AND (primary_ip = :h OR hostname = :h) LIMIT 1""",
            {"cid": case_id, "h": host},
        )
        if not asset:
            asset = fetchone(
                db,
                """INSERT INTO vuln_assets (case_id, hostname, primary_ip, asset_type, lifecycle_state)
                   VALUES (CAST(:cid AS uuid), :h, :h, 'host', 'active') RETURNING id""",
                {"cid": case_id, "h": host},
            )
        asset_id = str(asset["id"])
        vuln = {
            "plugin_name": name,
            "synopsis": synopsis or name,
            "description": description,
            "solution": solution,
            "severity": severity,
            "score": {"critical": 9.5, "high": 7.5, "medium": 5.0, "low": 2.0, "info": 0.0}.get(severity, 5.0),
            "port": port,
            "protocol": protocol,
            "scan_engine": "import",
        }
        if upsert_finding(
            db,
            case_id=case_id,
            scan_job_id=scan_job_id,
            asset_id=asset_id,
            vuln=vuln,
            source="import",
        ):
            imported += 1
            execute(
                db,
                """UPDATE vuln_findings SET description = :desc, scan_engine = 'import',
                   correlation_key = :ck, updated_at = NOW()
                   WHERE asset_id = CAST(:aid AS uuid) AND synopsis = :syn AND status = 'open'""",
                {
                    "desc": description,
                    "ck": correlation_key_for(host, port, name, None),
                    "aid": asset_id,
                    "syn": synopsis or name,
                },
            )
        refresh_asset_risk(db, asset_id)
    db.commit()
    return {"imported": imported, "case_id": case_id}
